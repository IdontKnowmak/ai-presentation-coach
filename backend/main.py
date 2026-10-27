from __future__ import annotations

import json
import logging
import os
import re
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import jwt
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.middleware.cors import CORSMiddleware
from pwdlib import PasswordHash
from pydantic import BaseModel, EmailStr, Field
from pydantic_settings import BaseSettings, SettingsConfigDict
from sqlalchemy import DateTime, ForeignKey, Integer, String, Text, create_engine, select, UniqueConstraint, inspect
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("presentation-coach")


class Settings(BaseSettings):
    # Resolve .env relative to this file, not the shell current working directory.
    model_config = SettingsConfigDict(
        env_file=Path(__file__).resolve().parent / ".env",
        env_file_encoding="utf-8",
        extra="ignore",
    )
    database_url: str = "sqlite:///./presentation_coach.db"
    cloudflare_account_id: str = ""
    cloudflare_api_token: str = ""
    cloudflare_model: str = "@cf/zai-org/glm-4.7-flash"
    cors_origins: str = "http://localhost:5173,http://127.0.0.1:5173"
    max_upload_mb: int = 20
    max_extracted_chars: int = 60000
    auth_secret_key: str = "CHANGE_ME_TO_A_LONG_RANDOM_SECRET_BEFORE_DEPLOYING"
    auth_token_minutes: int = 60
    auth_refresh_token_days: int = 30
    google_client_id: str = ""


settings = Settings()
password_hasher = PasswordHash.recommended()
bearer_scheme = HTTPBearer(auto_error=False)


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"
    __table_args__ = (UniqueConstraint("email", name="uq_users_email"),)
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    email: Mapped[str] = mapped_column(String(320), nullable=False, index=True)
    display_name: Mapped[str] = mapped_column(String(60), nullable=False, default="")
    password_hash: Mapped[str | None] = mapped_column(String(255), nullable=True)
    google_sub: Mapped[str | None] = mapped_column(String(255), nullable=True, unique=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    projects: Mapped[list["Project"]] = relationship(back_populates="owner", cascade="all, delete-orphan")


class Project(Base):
    __tablename__ = "projects"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    user_id: Mapped[int] = mapped_column(ForeignKey("users.id", ondelete="CASCADE"), index=True, nullable=False)
    owner: Mapped[User] = relationship(back_populates="projects")
    title: Mapped[str] = mapped_column(String(160), nullable=False)
    description: Mapped[str] = mapped_column(Text, default="")
    project_type: Mapped[str] = mapped_column(String(60), default="University")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc), onupdate=lambda: datetime.now(timezone.utc))
    presentations: Mapped[list["Presentation"]] = relationship(back_populates="project", cascade="all, delete-orphan")
    sessions: Mapped[list["PracticeSession"]] = relationship(back_populates="project", cascade="all, delete-orphan")


class Presentation(Base):
    __tablename__ = "presentations"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    project_id: Mapped[int] = mapped_column(ForeignKey("projects.id", ondelete="CASCADE"), index=True)
    filename: Mapped[str] = mapped_column(String(255), nullable=False)
    file_type: Mapped[str] = mapped_column(String(12), nullable=False)
    file_size: Mapped[int] = mapped_column(Integer, nullable=False)
    extracted_text: Mapped[str] = mapped_column(Text, default="")
    slide_count: Mapped[int] = mapped_column(Integer, default=0)
    # Legacy column retained for existing databases; new uploads never use Google Drive.
    drive_file_id: Mapped[str | None] = mapped_column(String(255), nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    project: Mapped[Project] = relationship(back_populates="presentations")
    questions: Mapped[list["PracticeQuestion"]] = relationship(back_populates="presentation", cascade="all, delete-orphan")


class PracticeQuestion(Base):
    __tablename__ = "practice_questions"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    presentation_id: Mapped[int] = mapped_column(ForeignKey("presentations.id", ondelete="CASCADE"), index=True)
    question_text: Mapped[str] = mapped_column(Text, nullable=False)
    category: Mapped[str] = mapped_column(String(80), default="CONTENT")
    difficulty: Mapped[str] = mapped_column(String(20), default="MEDIUM")
    slide_reference: Mapped[str] = mapped_column(String(120), default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    presentation: Mapped[Presentation] = relationship(back_populates="questions")


class PracticeSession(Base):
    __tablename__ = "practice_sessions"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    project_id: Mapped[int] = mapped_column(ForeignKey("projects.id", ondelete="CASCADE"), index=True)
    status: Mapped[str] = mapped_column(String(20), default="IN_PROGRESS")
    total_score: Mapped[int | None] = mapped_column(Integer, nullable=True)
    accuracy_score: Mapped[int | None] = mapped_column(Integer, nullable=True)
    completeness_score: Mapped[int | None] = mapped_column(Integer, nullable=True)
    reasoning_score: Mapped[int | None] = mapped_column(Integer, nullable=True)
    clarity_score: Mapped[int | None] = mapped_column(Integer, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    completed_at: Mapped[datetime | None] = mapped_column(DateTime(timezone=True), nullable=True)
    project: Mapped[Project] = relationship(back_populates="sessions")
    answers: Mapped[list["PracticeAnswer"]] = relationship(back_populates="session", cascade="all, delete-orphan")


class PracticeAnswer(Base):
    __tablename__ = "practice_answers"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    session_id: Mapped[int] = mapped_column(ForeignKey("practice_sessions.id", ondelete="CASCADE"), index=True)
    question_id: Mapped[int] = mapped_column(ForeignKey("practice_questions.id", ondelete="CASCADE"), index=True)
    answer_text: Mapped[str] = mapped_column(Text, nullable=False)
    accuracy_score: Mapped[int] = mapped_column(Integer, default=0)
    completeness_score: Mapped[int] = mapped_column(Integer, default=0)
    reasoning_score: Mapped[int] = mapped_column(Integer, default=0)
    clarity_score: Mapped[int] = mapped_column(Integer, default=0)
    total_score: Mapped[int] = mapped_column(Integer, default=0)
    strengths_json: Mapped[str] = mapped_column(Text, default="[]")
    improvements_json: Mapped[str] = mapped_column(Text, default="[]")
    model_answer: Mapped[str] = mapped_column(Text, default="")
    summary: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    session: Mapped[PracticeSession] = relationship(back_populates="answers")


def make_engine(url: str):
    # requirements.txt installs psycopg v3, so explicitly select its SQLAlchemy dialect.
    if url.startswith("postgres://"):
        url = "postgresql+psycopg://" + url.removeprefix("postgres://")
    elif url.startswith("postgresql://"):
        url = "postgresql+psycopg://" + url.removeprefix("postgresql://")
    connect_args = {"check_same_thread": False} if url.startswith("sqlite") else {}
    return create_engine(url, connect_args=connect_args, pool_pre_ping=True)


engine = make_engine(settings.database_url)
SessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


app = FastAPI(
    title="AI Presentation Coach API",
    version="0.1.0",
    description="Backend API for presentation projects, slide extraction, AI questions, and answer feedback.",
)
origins = [item.strip() for item in settings.cors_origins.split(",") if item.strip()]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=False, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def startup() -> None:
    Base.metadata.create_all(bind=engine)
    # Backward-compatible migration for databases created before account ownership existed.
    columns = {column["name"] for column in inspect(engine).get_columns("projects")}
    if "user_id" not in columns:
        with engine.begin() as connection:
            connection.exec_driver_sql('ALTER TABLE projects ADD COLUMN user_id INTEGER REFERENCES users(id)')
        logger.info("Added projects.user_id; legacy projects remain private/unassigned until manually migrated")
    if settings.auth_secret_key == "CHANGE_ME_TO_A_LONG_RANDOM_SECRET_BEFORE_DEPLOYING" or len(settings.auth_secret_key) < 32:
        logger.warning("AUTH_SECRET_KEY is missing or weak. Set a long random secret before deployment.")
    logger.info("Database tables initialized")


class RegisterRequest(BaseModel):
    name: str = Field(min_length=1, max_length=60)
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class LoginRequest(BaseModel):
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class RefreshRequest(BaseModel):
    refresh_token: str = Field(min_length=20, max_length=10000)


class GoogleLoginRequest(BaseModel):
    credential: str = Field(min_length=20, max_length=10000)


class ProjectCreate(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    description: str = Field(default="", max_length=5000)
    project_type: str = Field(default="University", max_length=60)


class ProjectUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=160)
    description: str | None = Field(default=None, max_length=5000)
    project_type: str | None = Field(default=None, max_length=60)


class QuestionGenerationRequest(BaseModel):
    count: int = Field(default=8, ge=3, le=20)
    language: str = Field(default="th", pattern="^(th|en)$")
    difficulty: str = Field(default="mixed", pattern="^(easy|medium|hard|mixed)$")


class SessionCreate(BaseModel):
    project_id: int


class EvaluationRequest(BaseModel):
    question_id: int
    answer: str = Field(min_length=1, max_length=12000)
    language: str = Field(default="th", pattern="^(th|en)$")


def user_dict(user: User) -> dict[str, Any]:
    return {"id": user.id, "email": user.email, "name": user.display_name}


def issue_token(user: User) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode({"sub": str(user.id), "token_use": "access", "iat": now, "exp": now + timedelta(minutes=settings.auth_token_minutes)}, settings.auth_secret_key, algorithm="HS256")


def issue_refresh_token(user: User) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode({"sub": str(user.id), "token_use": "refresh", "iat": now, "exp": now + timedelta(days=settings.auth_refresh_token_days)}, settings.auth_secret_key, algorithm="HS256")


def auth_tokens(user: User) -> dict[str, Any]:
    return {"access_token": issue_token(user), "refresh_token": issue_refresh_token(user), "token_type": "bearer", "user": user_dict(user)}


def current_user(credentials: HTTPAuthorizationCredentials | None = Depends(bearer_scheme), db: Session = Depends(get_db)) -> User:
    if not credentials:
        raise HTTPException(status_code=401, detail="Please sign in to continue.")
    try:
        payload = jwt.decode(credentials.credentials, settings.auth_secret_key, algorithms=["HS256"])
        if payload.get("token_use", "access") != "access":
            raise jwt.InvalidTokenError("Expected an access token")
        user_id = int(payload.get("sub", ""))
    except (jwt.PyJWTError, TypeError, ValueError):
        raise HTTPException(status_code=401, detail="Your session is invalid or expired. Please sign in again.")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=401, detail="Account not found. Please sign in again.")
    return user


def owned_project(project_id: int, user: User, db: Session) -> Project:
    project = db.scalar(select(Project).where(Project.id == project_id, Project.user_id == user.id))
    if not project:
        raise HTTPException(status_code=404, detail="Project not found.")
    return project


def owned_presentation(presentation_id: int, user: User, db: Session) -> Presentation:
    presentation = db.get(Presentation, presentation_id)
    if not presentation:
        raise HTTPException(status_code=404, detail="Presentation not found.")
    owned_project(presentation.project_id, user, db)
    return presentation


def owned_session(session_id: int, user: User, db: Session) -> PracticeSession:
    session = db.get(PracticeSession, session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Practice session not found.")
    owned_project(session.project_id, user, db)
    return session


def project_dict(project: Project) -> dict[str, Any]:
    presentations = sorted(project.presentations, key=lambda p: p.created_at, reverse=True)
    sessions = sorted(project.sessions, key=lambda s: s.created_at, reverse=True)
    completed = [s for s in sessions if s.status == "COMPLETED" and s.total_score is not None]
    return {
        "id": project.id, "title": project.title, "description": project.description, "type": project.project_type,
        "created_at": project.created_at.isoformat() if project.created_at else None,
        "updated_at": project.updated_at.isoformat() if project.updated_at else None,
        "presentations": [{"id": p.id, "filename": p.filename, "file_type": p.file_type, "file_size": p.file_size,
                           "slide_count": p.slide_count, "question_count": len(p.questions),
                           "created_at": p.created_at.isoformat() if p.created_at else None} for p in presentations],
        "questions": sum(len(p.questions) for p in presentations), "sessions": len(sessions),
        "score": round(sum(s.total_score for s in completed) / len(completed)) if completed else None,
    }


def extract_pdf(path: str) -> tuple[str, int]:
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page_number, page in enumerate(reader.pages, start=1):
        value = (page.extract_text() or "").strip()
        if value:
            parts.append(f"[Page {page_number}]\n{value}")
    return "\n\n".join(parts), len(reader.pages)


def extract_pptx(path: str) -> tuple[str, int]:
    from pptx import Presentation as PptxPresentation
    deck = PptxPresentation(path)
    parts = []
    for slide_number, slide in enumerate(deck.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                lines.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
        if lines:
            parts.append(f"[Slide {slide_number}]\n" + "\n".join(lines))
    return "\n\n".join(parts), len(deck.slides)


def generate_with_cloudflare(prompt: str) -> str:
    if not settings.cloudflare_account_id or not settings.cloudflare_api_token:
        raise HTTPException(
            status_code=503,
            detail="Cloudflare Workers AI is not configured. Add CLOUDFLARE_ACCOUNT_ID and CLOUDFLARE_API_TOKEN to backend/.env.",
        )
    try:
        from openai import OpenAI, RateLimitError
    except ImportError as exc:
        logger.exception("OpenAI SDK is missing from the backend environment")
        raise HTTPException(
            status_code=503,
            detail="The OpenAI SDK is not installed in the backend environment. Run: python -m pip install -r requirements.txt",
        ) from exc

    try:
        client = OpenAI(
            api_key=settings.cloudflare_api_token,
            base_url=(
                "https://api.cloudflare.com/client/v4/accounts/"
                f"{settings.cloudflare_account_id}/ai/v1"
            ),
            timeout=60.0,
            max_retries=1,
        )
        response = client.chat.completions.create(
            model=settings.cloudflare_model,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are an AI presentation coach. Follow the user's instructions, "
                        "ground answers only in the supplied slide content, do not invent facts, "
                        "match the requested language for all user-facing text, keep responses concise, "
                        "and return valid JSON only when JSON is requested."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0.2,
            max_completion_tokens=3000,
            extra_body={
                # GLM-4.7-Flash enables reasoning by default. Disable it for
                # this JSON/text task so the output budget is used for the answer.
                "chat_template_kwargs": {"enable_thinking": False}
            },
        )
        if not response.choices:
            logger.warning(
                "Cloudflare Workers AI returned no choices (model=%s, usage=%s)",
                settings.cloudflare_model,
                response.usage,
            )
            raise ValueError("Cloudflare Workers AI returned no completion choices.")

        choice = response.choices[0]
        result = choice.message.content
        if not result or not result.strip():
            logger.warning(
                "Cloudflare Workers AI returned empty content (model=%s, finish_reason=%s, usage=%s)",
                settings.cloudflare_model,
                choice.finish_reason,
                response.usage,
            )
            raise ValueError(
                "Cloudflare Workers AI returned empty text. Check the model output and finish reason."
            )
        return result
    except HTTPException:
        raise
    except RateLimitError as exc:
        logger.warning("Cloudflare Workers AI rate or quota limit reached")
        raise HTTPException(
            status_code=429,
            detail="Cloudflare AI rate or quota limit reached. Check Workers AI usage and try again later.",
        ) from exc
    except Exception as exc:
        logger.exception("Cloudflare Workers AI request failed")
        raise HTTPException(status_code=502, detail=f"AI request failed: {type(exc).__name__}") from exc


def parse_json_response(raw: str) -> Any:
    cleaned = re.sub(r"^\s*```(?:json)?\s*", "", raw.strip(), flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```\s*$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError as exc:
        logger.warning("AI returned invalid JSON")
        raise HTTPException(status_code=502, detail="AI returned an invalid structured response. Please try again.") from exc


@app.get("/")
def root():
    return {"name": "AI Presentation Coach API", "docs": "/docs", "health": "/health"}


@app.get("/health")
def health():
    return {
        "status": "ok",
        "database": "configured",
        "ai_provider": "cloudflare_workers_ai",
        "cloudflare_ai_configured": bool(settings.cloudflare_account_id and settings.cloudflare_api_token),
        "cloudflare_model": settings.cloudflare_model,
        "google_sign_in_configured": bool(settings.google_client_id),
    }


@app.post("/api/ai/test")
def test_cloudflare_ai(user: User = Depends(current_user)):
    """Verify Cloudflare Workers AI connectivity with a small, fixed prompt."""
    prompt = (
        'Return valid JSON only with two keys: "status" and "message". '
        'Set status to "ok" and message to "Cloudflare Workers AI connection successful".'
    )
    result = parse_json_response(generate_with_cloudflare(prompt))
    if not isinstance(result, dict) or result.get("status") != "ok":
        raise HTTPException(status_code=502, detail="Cloudflare AI responded, but the test response was unexpected.")
    return {"status": "ok", "provider": "cloudflare_workers_ai", "model": settings.cloudflare_model, "response": result}


@app.post("/api/auth/register", status_code=201)
def register(payload: RegisterRequest, db: Session = Depends(get_db)):
    email = str(payload.email).strip().lower()
    if db.scalar(select(User).where(User.email == email)):
        raise HTTPException(status_code=409, detail="An account with this email already exists.")
    user = User(email=email, display_name=payload.name.strip(), password_hash=password_hasher.hash(payload.password))
    db.add(user)
    db.commit()
    db.refresh(user)
    return auth_tokens(user)


@app.post("/api/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    email = str(payload.email).strip().lower()
    user = db.scalar(select(User).where(User.email == email))
    if not user or not user.password_hash or not password_hasher.verify(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    return auth_tokens(user)


@app.post("/api/auth/google")
def google_login(payload: GoogleLoginRequest, db: Session = Depends(get_db)):
    if not settings.google_client_id:
        raise HTTPException(status_code=503, detail="Google sign-in is not configured on the server.")
    try:
        from google.auth.transport import requests as google_requests
        from google.oauth2 import id_token
        claims = id_token.verify_oauth2_token(payload.credential, google_requests.Request(), settings.google_client_id)
        if not claims.get("email_verified") or not claims.get("sub") or not claims.get("email"):
            raise ValueError("Google account email is not verified")
    except Exception as exc:
        logger.info("Google token verification failed: %s", type(exc).__name__)
        raise HTTPException(status_code=401, detail="Could not verify your Google account. Please try again.") from exc
    email = str(claims["email"]).lower()
    user = db.scalar(select(User).where(User.google_sub == claims["sub"]))
    if not user:
        user = db.scalar(select(User).where(User.email == email))
        if user and user.google_sub and user.google_sub != claims["sub"]:
            raise HTTPException(status_code=409, detail="This email is linked to another Google account.")
        if not user:
            user = User(email=email, display_name=str(claims.get("name") or email.split("@")[0])[:60], google_sub=str(claims["sub"]))
            db.add(user)
        else:
            user.google_sub = str(claims["sub"])
            if not user.display_name:
                user.display_name = str(claims.get("name") or email.split("@")[0])[:60]
        db.commit()
        db.refresh(user)
    return auth_tokens(user)


@app.post("/api/auth/refresh")
def refresh_auth(payload: RefreshRequest, db: Session = Depends(get_db)):
    try:
        claims = jwt.decode(payload.refresh_token, settings.auth_secret_key, algorithms=["HS256"])
        if claims.get("token_use") != "refresh":
            raise jwt.InvalidTokenError("Expected a refresh token")
        user_id = int(claims.get("sub", ""))
    except (jwt.PyJWTError, TypeError, ValueError):
        raise HTTPException(status_code=401, detail="Your session has expired. Please sign in again.")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(status_code=401, detail="Account not found. Please sign in again.")
    return auth_tokens(user)


@app.get("/api/auth/me")
def get_me(user: User = Depends(current_user)):
    return user_dict(user)


@app.get("/api/projects")
def list_projects(user: User = Depends(current_user), db: Session = Depends(get_db)):
    return [project_dict(p) for p in db.scalars(select(Project).where(Project.user_id == user.id).order_by(Project.updated_at.desc())).all()]


@app.post("/api/projects", status_code=201)
def create_project(payload: ProjectCreate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = Project(user_id=user.id, title=payload.title.strip(), description=payload.description.strip(), project_type=payload.project_type)
    db.add(project)
    db.commit()
    db.refresh(project)
    return project_dict(project)


@app.get("/api/projects/{project_id}")
def get_project(project_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = owned_project(project_id, user, db)
    return project_dict(project)


@app.patch("/api/projects/{project_id}")
def update_project(project_id: int, payload: ProjectUpdate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = owned_project(project_id, user, db)
    for key, value in payload.model_dump(exclude_unset=True).items():
        if value is not None:
            setattr(project, key, value.strip() if isinstance(value, str) else value)
    project.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(project)
    return project_dict(project)


@app.delete("/api/projects/{project_id}", status_code=204)
def delete_project(project_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = owned_project(project_id, user, db)
    db.delete(project)
    db.commit()
    return None


@app.post("/api/projects/{project_id}/presentations", status_code=201)
async def upload_presentation(project_id: int, file: UploadFile = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = owned_project(project_id, user, db)
    filename = Path(file.filename or "presentation").name
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".pptx"}:
        raise HTTPException(status_code=415, detail="Only PDF and PPTX files are supported.")
    max_bytes = settings.max_upload_mb * 1024 * 1024
    content = await file.read(max_bytes + 1)
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"File exceeds the {settings.max_upload_mb} MB limit.")
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp:
            temp.write(content)
            temp_path = temp.name
        try:
            extracted, page_count = extract_pdf(temp_path) if suffix == ".pdf" else extract_pptx(temp_path)
        except Exception as exc:
            logger.exception("Could not extract uploaded document")
            raise HTTPException(status_code=422, detail="Could not read this document. Check that the PDF/PPTX is valid and not password-protected.") from exc
        # PDF text extractors can emit NUL bytes, which PostgreSQL text columns reject.
        extracted = extracted.replace(chr(0), "")
        if not extracted.strip():
            raise HTTPException(status_code=422, detail="No readable text was found. Scanned PDFs need OCR, which is not included yet.")
        extracted = extracted[:settings.max_extracted_chars]
        # The original upload exists only in a temporary file and is deleted in finally.
        # Persist extracted text and metadata in the database; never store the original file.
        presentation = Presentation(project_id=project_id, filename=filename, file_type=suffix.lstrip(".").upper(),
                                    file_size=len(content), extracted_text=extracted, slide_count=page_count)
        db.add(presentation)
        project.updated_at = datetime.now(timezone.utc)
        db.commit()
        db.refresh(presentation)
        return {"id": presentation.id, "project_id": project_id, "filename": filename, "file_type": presentation.file_type,
                "file_size": presentation.file_size, "slide_count": page_count, "extracted_characters": len(extracted),
                "original_file_stored": False, "message": "Presentation text extracted and saved to the database. Original file was not stored."}
    finally:
        await file.close()
        if temp_path and os.path.exists(temp_path):
            os.unlink(temp_path)


@app.get("/api/presentations/{presentation_id}")
def get_presentation(presentation_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = owned_presentation(presentation_id, user, db)
    return {"id": p.id, "project_id": p.project_id, "filename": p.filename, "file_type": p.file_type,
            "file_size": p.file_size, "slide_count": p.slide_count, "extracted_text": p.extracted_text,
            "questions": [{"id": q.id, "question": q.question_text, "category": q.category,
                           "difficulty": q.difficulty, "slide_reference": q.slide_reference} for q in p.questions]}


@app.post("/api/presentations/{presentation_id}/generate-questions")
def generate_questions(presentation_id: int, payload: QuestionGenerationRequest, user: User = Depends(current_user), db: Session = Depends(get_db)):
    p = owned_presentation(presentation_id, user, db)
    language = "Thai" if payload.language == "th" else "English"
    prompt = f"""You are a university presentation examiner. Generate exactly {payload.count} oral-defense questions grounded ONLY in the slide content below.
Language: {language}. Difficulty preference: {payload.difficulty}.
For Thai, write every question and category in natural Thai (keep only technical terms, product names, and acronyms in their usual form). For English, write them in English. Keep each question focused and concise, usually one sentence. Do not invent facts absent from the slides. Include conceptual, design-choice, limitation, evidence/testing, and application questions when relevant.
Return valid JSON only as an array of objects with keys: question (string), category (string), difficulty (EASY|MEDIUM|HARD), slide_reference (string). Keep JSON keys and difficulty values exactly as specified. Slide references must match labels present in the source, or be empty.
SLIDE CONTENT:
{p.extracted_text}"""
    data = parse_json_response(generate_with_cloudflare(prompt))
    if not isinstance(data, list) or not data:
        raise HTTPException(status_code=502, detail="AI did not return a list of questions.")
    saved = []
    for item in data[:payload.count]:
        if not isinstance(item, dict) or not str(item.get("question", "")).strip():
            continue
        question = PracticeQuestion(presentation_id=presentation_id, question_text=str(item["question"]).strip()[:2000],
                                    category=str(item.get("category", "CONTENT"))[:80],
                                    difficulty=str(item.get("difficulty", "MEDIUM")).upper()[:20],
                                    slide_reference=str(item.get("slide_reference", ""))[:120])
        db.add(question)
        saved.append(question)
    if not saved:
        db.rollback()
        raise HTTPException(status_code=502, detail="AI response contained no usable questions.")
    db.commit()
    for q in saved:
        db.refresh(q)
    return {"presentation_id": presentation_id, "count": len(saved), "questions": [
        {"id": q.id, "question": q.question_text, "category": q.category, "difficulty": q.difficulty, "slide_reference": q.slide_reference}
        for q in saved]}


@app.post("/api/sessions", status_code=201)
def create_session(payload: SessionCreate, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = owned_project(payload.project_id, user, db)
    session = PracticeSession(project_id=project.id)
    db.add(session)
    db.commit()
    db.refresh(session)
    return {"id": session.id, "project_id": session.project_id, "status": session.status, "created_at": session.created_at.isoformat()}


@app.post("/api/sessions/{session_id}/evaluate")
def evaluate_answer(session_id: int, payload: EvaluationRequest, user: User = Depends(current_user), db: Session = Depends(get_db)):
    session = owned_session(session_id, user, db)
    if not payload.answer.strip():
        raise HTTPException(status_code=422, detail="Please enter an answer before submitting.")
    if session.status == "COMPLETED":
        raise HTTPException(status_code=409, detail="This practice session is already completed.")
    question = db.get(PracticeQuestion, payload.question_id)
    if not question:
        raise HTTPException(status_code=404, detail="Question not found.")
    presentation = owned_presentation(question.presentation_id, user, db)
    if not presentation or presentation.project_id != session.project_id:
        raise HTTPException(status_code=400, detail="The question does not belong to this session's project.")

    # Idempotency: reuse an existing evaluation for this question/session instead of
    # calling the AI again when the client retries after a refresh or network error.
    existing = db.scalar(select(PracticeAnswer).where(
        PracticeAnswer.session_id == session.id,
        PracticeAnswer.question_id == question.id,
    ).order_by(PracticeAnswer.id.desc()))
    if existing and existing.answer_text.strip() == payload.answer.strip():
        return {
            "answer_id": existing.id, "session_id": session.id, "question_id": question.id,
            "scores": {
                "accuracy": existing.accuracy_score, "completeness": existing.completeness_score,
                "reasoning": existing.reasoning_score, "clarity": existing.clarity_score,
                "total": existing.total_score,
                "maxima": {"accuracy": 40, "completeness": 30, "reasoning": 20, "clarity": 10},
            },
            "strengths": json.loads(existing.strengths_json or "[]"),
            "improvements": json.loads(existing.improvements_json or "[]"),
            "model_answer": existing.model_answer, "summary": existing.summary,
            "notice": "Existing evaluation restored; AI was not called again.",
        }

    language = "Thai" if payload.language == "th" else "English"
    prompt = f"""Evaluate a student's oral-presentation answer using ONLY the provided slide content and question.
Language: {language}. All user-facing text in strengths, improvements, model_answer, and summary MUST be in this language. For Thai, use natural Thai and keep technical terms in their usual form. Keep feedback concise, specific, and actionable: usually 1-3 short items per list, a brief model answer, and a one- or two-sentence summary. Score criteria: accuracy 0-40, completeness 0-30, reasoning 0-20, clarity 0-10. total_score must equal their sum.
Strict scoring: empty, random, repetitive nonsense, unrelated text, or answers that do not address the question must receive 0-5 total points at most. Do not reward length or fluent wording by itself. Judge relevance and correctness against the slide content. Award points only for correct, relevant, supported content; use low scores when the answer lacks substance. Be fair and constructive; do not reward unsupported claims. If the slides do not contain enough information to verify a claim, say so rather than inventing facts.
Return valid JSON only with keys: accuracy (integer), completeness (integer), reasoning (integer), clarity (integer), total_score (integer), strengths (array of strings), improvements (array of strings), model_answer (string), summary (string). Keep JSON keys in English; do not add any other keys.
QUESTION: {question.question_text}
STUDENT ANSWER: {payload.answer.strip()}
SLIDE CONTENT:
{presentation.extracted_text}"""
    result = parse_json_response(generate_with_cloudflare(prompt))
    if not isinstance(result, dict):
        raise HTTPException(status_code=502, detail="AI did not return an evaluation object.")
    maxima = {"accuracy": 40, "completeness": 30, "reasoning": 20, "clarity": 10}
    scores = {}
    for key, maximum in maxima.items():
        try:
            scores[key] = max(0, min(maximum, int(result.get(key, 0))))
        except (TypeError, ValueError):
            scores[key] = 0
    total = sum(scores.values())
    strengths = result.get("strengths", []) if isinstance(result.get("strengths", []), list) else []
    improvements = result.get("improvements", []) if isinstance(result.get("improvements", []), list) else []
    model_answer = str(result.get("model_answer", ""))
    summary = str(result.get("summary", ""))
    saved_answer = PracticeAnswer(
        session_id=session.id, question_id=question.id, answer_text=payload.answer,
        accuracy_score=scores["accuracy"], completeness_score=scores["completeness"],
        reasoning_score=scores["reasoning"], clarity_score=scores["clarity"], total_score=total,
        strengths_json=json.dumps(strengths, ensure_ascii=False),
        improvements_json=json.dumps(improvements, ensure_ascii=False),
        model_answer=model_answer, summary=summary,
    )
    db.add(saved_answer)
    db.flush()
    answers = db.scalars(select(PracticeAnswer).where(PracticeAnswer.session_id == session.id)).all()
    session.accuracy_score = round(sum(a.accuracy_score for a in answers) / len(answers))
    session.completeness_score = round(sum(a.completeness_score for a in answers) / len(answers))
    session.reasoning_score = round(sum(a.reasoning_score for a in answers) / len(answers))
    session.clarity_score = round(sum(a.clarity_score for a in answers) / len(answers))
    session.total_score = round(sum(a.total_score for a in answers) / len(answers))
    db.commit()
    return {"answer_id": saved_answer.id, "session_id": session.id, "question_id": question.id,
            "scores": {**scores, "total": total, "maxima": maxima},
            "strengths": strengths, "improvements": improvements,
            "model_answer": model_answer, "summary": summary,
            "notice": "AI-generated feedback; verify against your instructor's expectations."}


@app.post("/api/sessions/{session_id}/complete")
def complete_session(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    session = owned_session(session_id, user, db)
    if not session.answers:
        raise HTTPException(status_code=400, detail="Answer at least one question before completing the session.")
    session.status = "COMPLETED"
    session.completed_at = datetime.now(timezone.utc)
    db.commit()
    return {"id": session.id, "status": session.status, "total_score": session.total_score,
            "answers_saved": len(session.answers), "completed_at": session.completed_at.isoformat()}


@app.get("/api/sessions/{session_id}/answers")
def list_session_answers(session_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    session = owned_session(session_id, user, db)
    return [{"id": a.id, "question_id": a.question_id, "answer": a.answer_text,
             "scores": {"accuracy": a.accuracy_score, "completeness": a.completeness_score,
                        "reasoning": a.reasoning_score, "clarity": a.clarity_score, "total": a.total_score},
             "strengths": json.loads(a.strengths_json or "[]"),
             "improvements": json.loads(a.improvements_json or "[]"),
             "model_answer": a.model_answer, "summary": a.summary,
             "created_at": a.created_at.isoformat() if a.created_at else None} for a in session.answers]


@app.get("/api/projects/{project_id}/sessions")
def list_sessions(project_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    owned_project(project_id, user, db)
    sessions = db.scalars(select(PracticeSession).where(PracticeSession.project_id == project_id).order_by(PracticeSession.created_at.desc())).all()
    return [{"id": s.id, "project_id": s.project_id, "status": s.status, "total_score": s.total_score,
             "scores": {"accuracy": s.accuracy_score, "completeness": s.completeness_score,
                        "reasoning": s.reasoning_score, "clarity": s.clarity_score},
             "created_at": s.created_at.isoformat() if s.created_at else None,
             "completed_at": s.completed_at.isoformat() if s.completed_at else None} for s in sessions]

@app.get("/api/projects/{project_id}/recommendations")
def project_recommendations(project_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Build coaching tips from saved scores and answer feedback; this endpoint never calls an AI model."""
    project = owned_project(project_id, user, db)
    completed_sessions = db.scalars(
        select(PracticeSession).where(
            PracticeSession.project_id == project_id,
            PracticeSession.status == "COMPLETED",
        ).order_by(PracticeSession.completed_at.desc())
    ).all()
    answers = [
        answer
        for session in completed_sessions
        for answer in session.answers
    ]
    recommendations: list[dict[str, str]] = []
    skill_specs = [
        ("accuracy", "ความถูกต้อง", 40, "ตรวจสอบข้อเท็จจริงและอ้างอิงข้อมูลจากสไลด์ก่อนตอบ"),
        ("completeness", "ความครบถ้วน", 30, "ตอบให้ครอบคลุมประเด็นหลัก แล้วเสริมรายละเอียดที่จำเป็น"),
        ("reasoning", "การให้เหตุผล", 20, "อธิบายเหตุผลเบื้องหลังข้อสรุป และยกหลักฐานหรือตัวอย่างประกอบ"),
        ("clarity", "ความชัดเจน", 10, "เริ่มจากใจความสำคัญ ใช้ประโยคกระชับ และเรียงคำตอบเป็นลำดับ"),
    ]
    if answers:
        for key, label, maximum, tip in skill_specs:
            values = [getattr(answer, f"{key}_score") for answer in answers]
            average = round(sum(values) / len(values))
            percentage = round(average / maximum * 100)
            if percentage < 75:
                recommendations.append({
                    "skill": label,
                    "score": f"{average}/{maximum}",
                    "priority": "high" if percentage < 60 else "medium",
                    "message": tip,
                })
        if not recommendations:
            recommendations.append({
                "skill": "จุดแข็งโดยรวม",
                "score": f"{round(sum(a.total_score for a in answers) / len(answers))}/100",
                "priority": "positive",
                "message": "คะแนนทักษะอยู่ในระดับดี รักษาโครงสร้างคำตอบและฝึกกับคำถามที่ซับซ้อนขึ้น",
            })
        improvements: dict[str, int] = {}
        for answer in answers:
            try:
                items = json.loads(answer.improvements_json or "[]")
            except (TypeError, json.JSONDecodeError):
                items = []
            if isinstance(items, list):
                for item in items:
                    text_value = str(item).strip()
                    if text_value:
                        improvements[text_value] = improvements.get(text_value, 0) + 1
        recurring = sorted(improvements.items(), key=lambda item: (-item[1], item[0]))[:3]
        for item, count in recurring:
            recommendations.append({
                "skill": "ข้อเสนอแนะที่พบจากผลประเมิน",
                "score": f"พบ {count} ครั้ง",
                "priority": "medium",
                "message": item[:500],
            })
    return {
        "project_id": project.id,
        "project_title": project.title,
        "completed_sessions": len(completed_sessions),
        "evaluated_answers": len(answers),
        "recommendations": recommendations,
        "source": "saved_evaluations",
    }

